from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ITEMS = [
    ("化简：A=(sinx+cosx)^4+(sinx-cosx)^4。", "展开并配凑得 A=2(sin^4x+6sin^2xcos^2x+cos^4x)=2[(sin^2x+cos^2x)^2+4sin^2xcos^2x]=2(1+sin^22x)。"),
    ("证明：(1-cos2x)/(1+cos2x)=tan^2x。", "由 1-cos2x=2sin^2x，1+cos2x=2cos^2x，故原式=sin^2x/cos^2x=tan^2x。"),
    ("化简：sin3x·cosx-cos3x·sinx。", "利用 sinAcosB-cosAsinB=sin(A-B)，原式=sin(3x-x)=sin2x。"),
    ("化简：sin5x·sinx+cos5x·cosx。", "利用 cos(A-B)=cosAcosB+sinAsinB，原式=cos(5x-x)=cos4x。"),
    ("化简：(sin2x+sin4x)/(cos2x+cos4x)。", "和差化积：分子=2sin3xcosx，分母=2cos3xcosx，故原式=tan3x。"),
    ("化简：(1+tanx)^2/(1+tan^2x)。", "原式=(1+tan^2x+2tanx)/(1+tan^2x)=1+2tanx/(1+tan^2x)=1+sin2x。"),
    ("化简：sinx/(1+cosx)+(1+cosx)/sinx。", "第一项=tan(x/2)，第二项=cot(x/2)，和为 tanu+cotu=2/sin2u，取 u=x/2 得 2cscx。"),
    ("化简：cosx/(1-sinx)+(1-sinx)/cosx。", "将第一项有理化：cosx/(1-sinx)=(1+sinx)/cosx，故原式=[(1+sinx)+(1-sinx)]/cosx=2secx。"),
    ("化简：sin4x/(sinx·cosx)。", "sin4x=2sin2xcos2x=4sinxcosxcos2x，故原式=4cos2x。"),
    ("求值：cos20°·cos40°·cos80°。", "利用积化公式 cosxcos2xcos4x=sin8x/(8sinx)，取 x=20° 得原式=sin160°/(8sin20°)=1/8。"),
    ("化简：(sinx+sin3x+sin5x)/(cosx+cos3x+cos5x)。", "分子=sin3x(2cos2x+1)，分母=cos3x(2cos2x+1)，故原式=tan3x。"),
    ("化简：[(1-cosx)(1-cos2x)(1-cos4x)]/[(1+cosx)(1+cos2x)(1+cos4x)]。", "每一项 (1-cosθ)/(1+cosθ)=tan^2(θ/2)，故原式=tan^2(x/2)·tan^2x·tan^22x。"),
    ("化简：cosx+cos3x+cos5x+cos7x。", "分组：(cosx+cos7x)+(cos3x+cos5x)=2cos4xcos3x+2cos4xcosx=2cos4x(2cos2xcosx)=4cos4xcos2xcosx。"),
    ("化简：sinx·sin3x+cosx·cos3x。", "利用 cos(A-B)=cosAcosB+sinAsinB，原式=cos(3x-x)=cos2x。"),
    ("化简：(tanx-sinx)/(tanx+sinx)。", "代入 tanx=sinx/cosx，原式=[sinx(1-cosx)/cosx]/[sinx(1+cosx)/cosx]=(1-cosx)/(1+cosx)=tan^2(x/2)。"),
    ("化简：(secx-cosx)/(secx+cosx)。", "原式=[(1/cosx)-cosx]/[(1/cosx)+cosx]=(1-cos^2x)/(1+cos^2x)=sin^2x/(1+cos^2x)。"),
    ("化简：(sinx+sin2x+sin3x)/(cosx+cos2x+cos3x)。", "分子=sin2x(2cosx+1)，分母=cos2x(2cosx+1)，故原式=tan2x。"),
    ("化简：sin6x/sin3x。", "sin6x=2sin3xcos3x，故原式=2cos3x。"),
    ("化简：(cos2x-cos4x)/(sin4x-sin2x)。", "分子=2sin3xsinx，分母=2cos3xsinx，故原式=tan3x。"),
    ("化简：sinx·cos2x+cosx·sin2x。", "利用 sinAcosB+cosAsinB=sin(A+B)，原式=sin3x。"),
    ("已知 tanx+cotx=4，求 sin2x。", "tanx+cotx=(sin^2x+cos^2x)/(sinxcosx)=1/(sinxcosx)=2/sin2x=4，故 sin2x=1/2。"),
    ("已知 sinx-cosx=1/3，求 sin2x。", "(sinx-cosx)^2=1-2sinxcosx=1/9，故 sin2x=2sinxcosx=8/9。"),
    ("已知 sin2x=3/5 且 x∈(0,π/4)，求 tanx。", "设 t=tanx，2t/(1+t^2)=3/5，解得 3t^2-10t+3=0，t=3 或 1/3；由区间取 t=1/3。"),
    ("已知 cosx-sinx=1/2 且 x∈(0,π/2)，求 sin2x。", "(cosx-sinx)^2=1-2sinxcosx=1/4，故 sin2x=3/4。"),
    ("已知 secx+tanx=3，求 secx-tanx 与 sinx。", "由 (secx+tanx)(secx-tanx)=1，得 secx-tanx=1/3。联立得 secx=(3+1/3)/2=5/3，tanx=(3-1/3)/2=4/3，故 sinx=4/5。"),
    ("已知 tan(x/2)=2，求 sinx、cosx。", "半角代换：sinx=2t/(1+t^2)=4/5，cosx=(1-t^2)/(1+t^2)=-3/5。"),
    ("已知 sinx+cosx=1 且 x∈(0,π)，求 x。", "由 √2sin(x+π/4)=1，得 sin(x+π/4)=√2/2。解得 x=0 或 π/2；区间内仅 x=π/2。"),
    ("已知 sinx:cosx=1:2，x 为锐角，求 sin3x。", "sinx=1/√5，cosx=2/√5。sin3x=3sinx-4sin^3x=11/(5√5)=11√5/25。"),
    ("已知 cos2x=-3/5 且 x∈(0,π/2)，求 tanx。", "tan^2x=(1-cos2x)/(1+cos2x)=(1+3/5)/(1-3/5)=4，且锐角取正，tanx=2。"),
    ("已知 sinxcosx=1/4，x 为锐角，求 sinx+cosx。", "(sinx+cosx)^2=1+2sinxcosx=3/2，故 sinx+cosx=√(3/2)=√6/2。"),
    ("已知 x∈(0,π/4)，sinx+cosx=a(1<a<√2)，求 sin2x、cos2x。", "sin2x=(sinx+cosx)^2-1=a^2-1。又 2x∈(0,π/2) 故 cos2x>0，cos2x=√(1-(a^2-1)^2)。"),
    ("已知 tanx=2，求 tan2x 与 tan3x。", "tan2x=2tanx/(1-tan^2x)=4/(1-4)=-4/3。tan3x=(3t-t^3)/(1-3t^2)=2/11。"),
    ("已知 tanx=1/3，求 sin2x、cos2x。", "sin2x=2t/(1+t^2)=3/5，cos2x=(1-t^2)/(1+t^2)=4/5。"),
    ("已知 sin(x+y)=3/5 且 cos(x+y)>0，cos(x-y)=4/5，求 cos2x+cos2y。", "cos2x+cos2y=2cos(x+y)cos(x-y)。由 sin(x+y)=3/5 且 cos(x+y)>0 得 cos(x+y)=4/5，故结果=2·4/5·4/5=32/25。"),
    ("已知 sin(x+y)=a，sin(x-y)=b，且 cos(x+y)>0，求 sin2x-sin2y。", "sin2x-sin2y=2cos(x+y)sin(x-y)=2b√(1-a^2)。"),
    ("解方程：sin2x=sinx，x∈[0,2π)。", "sinx(2cosx-1)=0。故 x=0,π 或 cosx=1/2 给 x=π/3,5π/3。"),
    ("解方程：cos2x+sinx=0，x∈[0,2π)。", "1-2sin^2x+sinx=0，即 2s^2-s-1=0。s=1 或 -1/2。故 x=π/2,7π/6,11π/6。"),
    ("解方程：tan2x=√3，x∈(0,π)。", "2x=π/3+kπ，故 x=π/6+kπ/2。区间内解：x=π/6,2π/3。"),
    ("解方程：sinx+sin2x+sin3x=0，x∈(0,2π)。", "化为 sin2x(2cosx+1)=0。得 sin2x=0 或 cosx=-1/2。故 x=π/2,π,3π/2,2π/3,4π/3。"),
    ("解方程：cosx+cos2x+cos3x=0，x∈(0,2π)。", "化为 cos2x(2cosx+1)=0。得 cos2x=0 或 cosx=-1/2。故 x=π/4,3π/4,5π/4,7π/4,2π/3,4π/3。"),
    ("求函数 y=sinx+cosx+sinxcosx 的最大值。", "设 t=sinx+cosx∈[-√2,√2]，则 sinxcosx=(t^2-1)/2。故 y=t+(t^2-1)/2=0.5t^2+t-0.5，为开口向上二次函数，最大值在端点 t=√2 处取到，y_max=√2+1/2。"),
    ("求函数 y=(sinx+cosx)^2+1/(sinx+cosx)^2 的最小值。", "设 u=(sinx+cosx)^2>0，则 y=u+1/u≥2。当 u=1 时取等号，故最小值为 2。"),
    ("求函数 y=(1-sin2x)/(1+sin2x) 的值域。", "原式=[(sinx-cosx)^2]/[(sinx+cosx)^2]=tan^2(x-π/4)，故 y≥0 且可任意大，值域为 [0,+∞)。"),
    ("已知 y=sinx+cosx，求 F=y^4-4y^2 的最小值。", "y∈[-√2,√2]。令 f(t)=t^4-4t^2，考察区间端点与临界点 t=0,±√2。f(0)=0，f(±√2)=-4，故最小值为 -4。"),
    ("求函数 y=sin2x+cos2x 的最大值。", "y=√2sin(2x+π/4)，故最大值为 √2。"),
    ("设 a=sinx+cosx，b=sinx-cosx。求 a^2+b^2 与 sin2x 的 a,b 表达式。", "a^2+b^2=2(sin^2x+cos^2x)=2；a^2-b^2=4sinxcosx=2sin2x，故 sin2x=(a^2-b^2)/2。"),
    ("证明：(tanx+cotx)^2=(secx·cscx)^2。", "左边=[(sin^2x+cos^2x)/(sinxcosx)]^2=1/(sin^2xcos^2x)=sec^2x·csc^2x，即右边。"),
    ("化简：(sin3x+sinx)/(cos3x-cosx)。", "分子=2sin2xcosx，分母=-2sin2xsinx，故原式=-cotx。"),
    ("化简：(cos5x-cosx)/(sin5x+sinx)。", "分子=-2sin3xsin2x，分母=2sin3xcos2x，故原式=-tan2x。"),
    ("已知 x∈(0,π/2)，sinx+cosx=√5/2，求 tanx+cotx 与 cos4x。", "sin2x=(sinx+cosx)^2-1=1/4。故 tanx+cotx=2/sin2x=8；cos4x=1-2sin^22x=1-2·(1/16)=7/8。"),
]


def apply_chinese_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(11)


def build_exercise_book(path: str) -> None:
    doc = Document()
    apply_chinese_font(doc)
    doc.add_heading("高考130分以上难度：三角函数恒等变换习题册（50题）", level=0)
    doc.add_paragraph("说明：本套题聚焦恒等变换、条件求值与综合方程，难度对标高考130分以上冲刺。")
    doc.add_paragraph("建议：先独立完成，再对照答案册进行反思总结。")

    for i, (question, _) in enumerate(ITEMS, start=1):
        doc.add_paragraph(f"{i}. {question}")

    doc.save(path)


def build_answer_book(path: str) -> None:
    doc = Document()
    apply_chinese_font(doc)
    doc.add_heading("高考130分以上难度：三角函数恒等变换答案册（50题）", level=0)
    doc.add_paragraph("说明：答案强调关键变换路径，适合二轮复盘与错题精练。")

    for i, (question, answer) in enumerate(ITEMS, start=1):
        doc.add_heading(f"第{i}题", level=2)
        doc.add_paragraph(f"题目：{question}")
        doc.add_paragraph(f"答案：{answer}")

    doc.save(path)


if __name__ == "__main__":
    build_exercise_book("/workspace/高考130分以上难度_三角函数恒等变换习题册_50题.docx")
    build_answer_book("/workspace/高考130分以上难度_三角函数恒等变换答案册_50题.docx")
    print("Word 文件已生成。")
